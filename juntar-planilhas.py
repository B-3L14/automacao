import numpy as np
import pandas as pd
from pathlib import Path
from urllib.parse import urlparse, parse_qs     # NOVO: para ler o código do anúncio dentro do link
from openpyxl import load_workbook              # NOVO: para ler os links escondidos nas células
from openpyxl.styles import Font

# ======================================================================
# 1) CONFIGURAÇÕES  (só mexa aqui se mudar o nome ou a pasta dos arquivos)
# ======================================================================
PASTA = Path(".")  # "." = mesma pasta onde este script está salvo

ARQ_COM_VIDEO = PASTA / "../planilhas/inventario_drive.xlsx"            # Planilha A: com vídeo
ARQ_SEM_VIDEO = PASTA / "../planilhas/produtos_3_fora_do_drive.xlsx"    # Planilha B: sem vídeo
ARQ_ANUNCIOS = PASTA / "../planilhas/Resumo_Criterios__1_.xlsx"         # Planilha C: anúncios
ARQ_SAIDA = PASTA / "../planilhas/produtos_video_anuncio.xlsx"          # Arquivo que será criado

# EAN que está digitado errado na planilha de anúncios -> EAN correto do Drive.
# (Usado só para comparar. Na planilha final continua o EAN original do anúncio.)
CORRECOES_EAN = {"78999473100816": "7899947310081"}

# Marcas escritas de forma diferente no Drive -> nome usado nos anúncios.
CORRECOES_MARCA = {"ACTE": "ACTE SPORTS", "VOLLO": "VOLLO SPORTS"}

# Colunas da planilha final, na ordem pedida.
COLUNAS_FINAIS = [
    "SKU", "EAN", "MLB", "Marca", "Título",
    "Possui anúncio", "Possui vídeo",
    "Tem 1 clipe?",                                         # NOVO
]


# ======================================================================
# 2) FUNÇÕES AUXILIARES (pequenas tarefas de limpeza)
# ======================================================================
def limpar_texto(serie):
    """Transforma em texto, tira espaços das pontas e o '.0' que o Excel às vezes coloca."""
    return serie.astype("string").str.strip().str.replace(r"\.0$", "", regex=True)


def preparar_marca(serie):
    """Marca em maiúsculas e com os nomes padronizados."""
    return limpar_texto(serie).str.upper().replace(CORRECOES_MARCA)


def criar_chave_ean(serie_ean):
    """
    Cria a 'chave' usada para saber se dois produtos são o mesmo:
    aplica a correção de EAN e tira os zeros do começo.
    """
    return serie_ean.replace(CORRECOES_EAN).str.lstrip("0")


def ler_aba(caminho, aba, linha_cabecalho):
    """Lê uma aba como texto (para o EAN não virar 7,89E+12) e remove linhas totalmente vazias."""
    df = pd.read_excel(caminho, sheet_name=aba, header=linha_cabecalho, dtype=str)
    return df.dropna(how="all")


# NOVO ------------------------------------------------------------------
def ler_links_clipe(caminho):
    """
    Lê os links que estão escondidos nas células da coluna "Tem 1 clipe?"
    da planilha de anúncios (o pandas não enxerga links, por isso uso o openpyxl).

    Devolve duas coisas:
      - links: dicionário {MLB: link}, só com os links confiáveis
      - errados: lista de MLBs cujo link aponta para OUTRO anúncio
    """
    planilha = load_workbook(caminho)
    aba = planilha.worksheets[0]

    # Descobre em qual coluna estão o MLB e o "Tem 1 clipe?" lendo o cabeçalho (1ª linha).
    cabecalho = [celula.value for celula in aba[1]]
    col_mlb = cabecalho.index("MLB") + 1
    col_clipe = cabecalho.index("Tem 1 clipe?") + 1

    links, errados = {}, []
    for linha in range(2, aba.max_row + 1):
        mlb = aba.cell(linha, col_mlb).value
        link = aba.cell(linha, col_clipe).hyperlink
        if mlb is None or link is None or not link.target:
            continue

        mlb = str(mlb).strip()
        # O link traz o código do anúncio no trecho "item_id=MLB...".
        # Só aceito o link se esse código for o mesmo MLB da linha.
        item_do_link = parse_qs(urlparse(link.target).query).get("item_id", [""])[0]
        if item_do_link == mlb:
            links[mlb] = link.target
        else:
            errados.append(mlb)
    return links, errados
# -----------------------------------------------------------------------


# ======================================================================
# 3) MONTAGEM DA TABELA
# ======================================================================
def montar_tabela():
    avisos = []

    # ---------- Planilha A: produtos com vídeo no Drive ----------
    # (o cabeçalho está na 2ª linha da aba "Geral", por isso linha_cabecalho=1)
    a = ler_aba(ARQ_COM_VIDEO, "Geral", 1)
    a["EAN"] = limpar_texto(a["EAN"])
    a["Marca"] = preparar_marca(a["Marca"])
    a["chave"] = criar_chave_ean(a["EAN"])
    # Um produto só "tem vídeo" se existir pelo menos um arquivo .mp4 nele.
    a["e_mp4"] = a["Nome do Arquivo"].fillna("").str.lower().str.endswith(".mp4")

    # Cada produto aparece em várias linhas (um arquivo por linha): junta em 1 linha por produto.
    a_prod = a.groupby("chave", as_index=False).agg(
        EAN=("EAN", "first"), Marca=("Marca", "first"), tem_video=("e_mp4", "any")
    )
    chaves_com_video = set(a_prod.loc[a_prod["tem_video"], "chave"])

    pastas_vazias = a_prod.loc[~a_prod["tem_video"], "EAN"].tolist()
    if pastas_vazias:
        avisos.append(f"EANs com pasta no Drive mas SEM nenhum .mp4 (marcados como 'Não'): {pastas_vazias}")

    # ---------- Planilha B: produtos sem vídeo ----------
    b = ler_aba(ARQ_SEM_VIDEO, "Geral", 1)
    b["EAN"] = limpar_texto(b["EAN"])
    b["Marca"] = preparar_marca(b["Marca"])
    b["chave"] = criar_chave_ean(b["EAN"])
    b_prod = b.drop_duplicates("chave")[["chave", "EAN", "Marca"]]

    conflito = chaves_com_video & set(b_prod["chave"])
    if conflito:
        avisos.append(f"EANs que aparecem COM e SEM vídeo ao mesmo tempo: {sorted(conflito)}")

    # ---------- Planilha C: anúncios (a primeira aba, cabeçalho na 1ª linha) ----------
    c = ler_aba(ARQ_ANUNCIOS, 0, 0)
    for coluna in ["SKU", "EAN", "MLB"]:
        c[coluna] = limpar_texto(c[coluna])
    c["Título"] = c["Título"].astype("string").str.strip()
    c["Marca"] = preparar_marca(c["Marca"])
    c["chave"] = criar_chave_ean(c["EAN"])
    chaves_com_anuncio = set(c["chave"])

    # NOVO: "Tem 1 clipe?" e o link de cada anúncio ------------------------
    c["Tem 1 clipe?"] = limpar_texto(c["Tem 1 clipe?"])
    links_clipe, mlbs_link_errado = ler_links_clipe(ARQ_ANUNCIOS)
    c["Link do clipe"] = c["MLB"].map(links_clipe)   # anúncio sem link confiável fica vazio
    if mlbs_link_errado:
        avisos.append(
            "Anúncios cujo link do clipe, na planilha de origem, aponta para OUTRO anúncio "
            f"(ficaram SEM link): {mlbs_link_errado}"
        )
    # ------------------------------------------------------------------------

    corrigidos = c.loc[c["EAN"].isin(CORRECOES_EAN.keys()), "EAN"].unique().tolist()
    if corrigidos:
        avisos.append(f"EAN(s) de anúncio corrigido(s) só para comparação: {corrigidos}")

    todas_chaves_drive = chaves_com_video | set(a_prod["chave"]) | set(b_prod["chave"])
    sem_drive = sorted(set(c.loc[~c["chave"].isin(todas_chaves_drive), "EAN"]))
    if sem_drive:
        avisos.append(f"EANs de anúncio que não existem nas planilhas de vídeo: {sem_drive}")

    # ---------- Lista única de produtos (sem repetir) ----------
    # Para EAN e Marca, a ordem de preferência é: anúncios, depois sem vídeo, depois com vídeo.
    c_produto = c.drop_duplicates("chave")[["chave", "EAN", "Marca"]]
    produtos = pd.concat([c_produto, b_prod, a_prod[["chave", "EAN", "Marca"]]], ignore_index=True)
    produtos = produtos.drop_duplicates("chave", keep="first")

    # ---------- Colunas "Possui vídeo" e "Possui anúncio" ----------
    produtos["Possui vídeo"] = np.where(produtos["chave"].isin(chaves_com_video), "Sim", "Não")
    produtos["Possui anúncio"] = np.where(produtos["chave"].isin(chaves_com_anuncio), "Sim", "Não")

    # ---------- Traz SKU, MLB, Título e clipe (uma linha por anúncio) ----------
    tabela = produtos.merge(
        c[["chave", "SKU", "MLB", "Título", "Tem 1 clipe?", "Link do clipe"]],   # NOVO: clipe e link
        on="chave",
        how="left",
    )

    # ---------- Ordenação em 4 grupos ----------
    tem_video = tabela["Possui vídeo"] == "Sim"
    tem_anuncio = tabela["Possui anúncio"] == "Sim"
    tabela["_grupo"] = np.select(
        [tem_video & tem_anuncio, tem_anuncio & ~tem_video, tem_video & ~tem_anuncio],
        [1, 2, 3],
        default=4,
    )
    tabela = tabela.sort_values(
        ["_grupo", "Marca", "EAN", "Título", "MLB"], na_position="last", kind="stable"
    ).reset_index(drop=True)

    return tabela, avisos


# ======================================================================
# 4) SALVAR EM EXCEL
# ======================================================================
def salvar_planilha(tabela, destino):
    saida = tabela[COLUNAS_FINAIS]
    links = tabela["Link do clipe"].tolist()                          # NOVO
    coluna_clipe = COLUNAS_FINAIS.index("Tem 1 clipe?") + 1           # NOVO: número da coluna (1 = A)

    with pd.ExcelWriter(destino, engine="openpyxl") as escritor:
        saida.to_excel(escritor, index=False, sheet_name="Produtos")
        aba = escritor.sheets["Produtos"]

        aba.freeze_panes = "A2"                 # cabeçalho fixo ao rolar
        aba.auto_filter.ref = aba.dimensions    # filtros nas colunas
        for celula in aba[1]:
            celula.font = Font(bold=True)       # cabeçalho em negrito

        larguras = {"A": 24, "B": 18, "C": 16, "D": 26, "E": 70, "F": 16, "G": 14, "H": 16}  # NOVO: coluna H
        for coluna, largura in larguras.items():
            aba.column_dimensions[coluna].width = largura

        # SKU, EAN e MLB como texto (não perde zeros à esquerda se você editar)
        for linha in aba.iter_rows(min_row=2, max_col=3):
            for celula in linha:
                celula.number_format = "@"

        # NOVO: coloca o link na célula "Tem 1 clipe?" (azul e sublinhado, como um link normal)
        for numero_linha, link in enumerate(links, start=2):   # linha 2 = primeiro produto
            if pd.notna(link):
                celula = aba.cell(row=numero_linha, column=coluna_clipe)
                celula.hyperlink = link
                celula.font = Font(color="0563C1", underline="single")


# ======================================================================
# 5) EXECUÇÃO
# ======================================================================
if __name__ == "__main__":
    tabela, avisos = montar_tabela()
    salvar_planilha(tabela, ARQ_SAIDA)

    nomes_grupos = {
        1: "1) Vídeo + anúncio",
        2: "2) Anúncio, sem vídeo",
        3: "3) Vídeo, sem anúncio",
        4: "4) Nenhum dos dois",
    }
    print(f"Arquivo salvo: {ARQ_SAIDA}  ({len(tabela)} linhas)\n")
    for numero, nome in nomes_grupos.items():
        parte = tabela[tabela["_grupo"] == numero]
        print(f"{nome}: {parte['EAN'].nunique()} produtos ({len(parte)} linhas)")

    if avisos:
        print("\nAvisos:")
        for aviso in avisos:
            print(" -", aviso)