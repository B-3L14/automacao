"""
planilha.py - le a planilha da tela [0667] e entrega a lista de produtos, na ordem.
"""

import unicodedata
from dataclasses import dataclass

from openpyxl import load_workbook

import config
from erros import ErroConfiguracao, ErroProdutoDiferente


@dataclass(frozen=True)
class Produto:
    codigo_interno: str
    sku: str          # Codigo Auxiliar
    descricao: str

    @property
    def sku_afiliado(self):
        """O 1SKU: o mesmo codigo com '1' na frente."""
        return config.PREFIXO_1SKU + self.sku


def _norm(valor):
    return unicodedata.normalize("NFC", str(valor)).strip().casefold()


def _texto(linha, indice):
    if indice >= len(linha) or linha[indice] is None:
        return ""
    valor = linha[indice]
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _checar_duplicados(produtos, campo, rotulo):
    vistos = set()
    for produto in produtos:
        valor = getattr(produto, campo)
        if valor in vistos:
            raise ErroConfiguracao(f"{rotulo} repetido na planilha: {valor}")
        vistos.add(valor)


def carregar_produtos(caminho):
    if not caminho.exists():
        raise ErroConfiguracao(f"planilha nao encontrada: {caminho}")

    workbook = load_workbook(caminho, read_only=True, data_only=True)
    try:
        linhas = workbook.worksheets[0].iter_rows(values_only=True)
        cabecalho = next(linhas, None)
        if cabecalho is None:
            raise ErroConfiguracao("a planilha esta vazia")

        colunas = {_norm(nome): i for i, nome in enumerate(cabecalho) if nome is not None}
        posicoes = {}
        for nome in (config.COL_CODIGO_INTERNO, config.COL_SKU, config.COL_DESCRICAO):
            if _norm(nome) not in colunas:
                raise ErroConfiguracao(f"coluna '{nome}' nao encontrada na planilha")
            posicoes[nome] = colunas[_norm(nome)]

        produtos = []
        for linha in linhas:
            sku = _texto(linha, posicoes[config.COL_SKU])
            if not sku:
                continue  # rodape (so tem a contagem) ou linha vazia

            codigo_interno = _texto(linha, posicoes[config.COL_CODIGO_INTERNO])
            if not codigo_interno:
                raise ErroConfiguracao(f"produto com SKU {sku} esta sem Codigo Interno na planilha")

            produtos.append(
                Produto(
                    codigo_interno=codigo_interno,
                    sku=sku,
                    descricao=_texto(linha, posicoes[config.COL_DESCRICAO]),
                )
            )
    finally:
        workbook.close()

    if not produtos:
        raise ErroConfiguracao("nenhum produto encontrado na planilha")

    _checar_duplicados(produtos, "codigo_interno", "Codigo Interno")
    _checar_duplicados(produtos, "sku", "Codigo Auxiliar")
    return produtos


def posicao_do_codigo(produtos, codigo_interno):
    for posicao, produto in enumerate(produtos):
        if produto.codigo_interno == codigo_interno:
            return posicao
    raise ErroProdutoDiferente(
        f"o produto selecionado na tela (Codigo Interno {codigo_interno}) nao esta na planilha"
    )